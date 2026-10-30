"""정보분석 경진대회 심사 시스템 — 데이터 저장, 진행 단계, 점수 집계.

Streamlit에 의존하지 않는 순수 로직이라 단독으로 테스트할 수 있습니다.
데이터는 SQLite 파일 하나(data/event.db)에 저장됩니다.
"""
from __future__ import annotations

import copy
import io
import json
import os
import secrets
import sqlite3
import threading
import time
from contextlib import contextmanager
from statistics import mean

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("JUDGING_DB", os.path.join(BASE, "data", "event.db"))
SETTINGS_FILE = os.path.join(BASE, "settings.json")
ONLINE_SEC = 15          # 이 시간 안에 신호가 있으면 '접속 중'
_lock = threading.RLock()

DEFAULT_STATE = {
    "phase": "lobby",        # lobby → order → running → done → reveal
    "order": [],             # 발표 순서 (팀 번호 목록)
    "drawn_at": 0.0,
    "cur": 0,                # 현재 발표 순번 (order의 인덱스)
    "step": "presenting",    # presenting(발표 중) / scoring(평가 요청)
    "closed": [],            # 평가가 마감된 팀 번호
    "revealed_at": 0.0,
    "winner_override": None, # 동점일 때 관리자가 정한 최우수 팀
}


class FlowError(ValueError):
    """진행 순서상 지금 할 수 없는 동작."""


# ───────────────────────── 저장소
def _connect() -> sqlite3.Connection:
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    c = sqlite3.connect(DB_PATH, timeout=20, isolation_level=None, check_same_thread=False)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA journal_mode=WAL")
    c.execute("PRAGMA busy_timeout=20000")
    return c


@contextmanager
def _tx():
    with _lock:
        c = _connect()
        try:
            c.execute("BEGIN IMMEDIATE")
            yield c
            c.execute("COMMIT")
        except BaseException:
            c.execute("ROLLBACK")
            raise
        finally:
            c.close()


def _get(c, key, default=None):
    r = c.execute("SELECT v FROM kv WHERE k=?", (key,)).fetchone()
    return json.loads(r["v"]) if r else default


def _put(c, key, value):
    c.execute("INSERT INTO kv(k, v) VALUES(?, ?) ON CONFLICT(k) DO UPDATE SET v=excluded.v",
              (key, json.dumps(value, ensure_ascii=False)))


def _make_seats(c, n):
    c.execute("DELETE FROM seats")
    used = set()
    for i in range(n):
        tok = secrets.token_hex(4)
        while tok in used:
            tok = secrets.token_hex(4)
        used.add(tok)
        c.execute("INSERT INTO seats(idx, token, claimed_at, last_seen) VALUES(?, ?, NULL, NULL)", (i, tok))


def load_file_settings() -> dict:
    with open(SETTINGS_FILE, encoding="utf-8") as f:
        return normalize_settings(json.load(f))


def init():
    """테이블을 만들고, 처음이면 settings.json으로 초기값을 채웁니다. 매 실행마다 호출해도 안전합니다."""
    with _tx() as c:
        c.execute("CREATE TABLE IF NOT EXISTS kv (k TEXT PRIMARY KEY, v TEXT NOT NULL)")
        c.execute("CREATE TABLE IF NOT EXISTS seats (idx INTEGER PRIMARY KEY, token TEXT UNIQUE NOT NULL,"
                  " claimed_at REAL, last_seen REAL)")
        c.execute("CREATE TABLE IF NOT EXISTS scores (token TEXT NOT NULL, team INTEGER NOT NULL, vals TEXT NOT NULL,"
                  " total REAL NOT NULL, proxy INTEGER NOT NULL DEFAULT 0, updated_at REAL NOT NULL,"
                  " PRIMARY KEY (token, team))")
        if _get(c, "settings") is None:
            s = load_file_settings()
            _put(c, "settings", s)
            _put(c, "state", copy.deepcopy(DEFAULT_STATE))
            _make_seats(c, len(s["judges"]))


# ───────────────────────── 설정
def _num(x, default=0.0):
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def normalize_settings(s: dict) -> dict:
    s = copy.deepcopy(s)
    s.setdefault("title_top", "")
    s.setdefault("title_main", "경진대회")
    s.setdefault("date_line", "")
    s["admin_pin"] = str(s.get("admin_pin", "1103")).strip() or "1103"
    s["judge_code"] = str(s.get("judge_code", "") or "").strip()
    s["public_url"] = str(s.get("public_url", "") or "").strip().rstrip("/")
    s["input_mode"] = "number" if s.get("input_mode") == "number" else "grade"
    labels = [str(x) for x in s.get("grade_labels", ["매우우수", "우수", "보통", "미흡", "매우미흡"])]
    ratio = [_num(x) for x in s.get("grade_ratio", [1.0, 0.8, 0.6, 0.4, 0.2])]
    k = min(len(labels), len(ratio)) or 5
    s["grade_labels"], s["grade_ratio"] = (labels[:k], ratio[:k]) if labels and ratio else (
        ["매우우수", "우수", "보통", "미흡", "매우미흡"], [1.0, 0.8, 0.6, 0.4, 0.2])
    s["trim_high_low"] = bool(s.get("trim_high_low", True))
    s["show_runner_ups"] = bool(s.get("show_runner_ups", False))
    s["guide_note"] = str(s.get("guide_note", "") or "").strip()
    s["guide"] = [{"name": str(g.get("name", "")).strip(), "weight": str(g.get("weight", "") or "").strip(),
                   "items": [str(x).strip() for x in (g.get("items") or []) if str(x).strip()]}
                  for g in (s.get("guide") or []) if isinstance(g, dict) and str(g.get("name", "")).strip()]
    s["teams"] = [{"name": str(t.get("name", "")).strip() or f"{i + 1}팀",
                   "dept": str(t.get("dept", "")).strip(),
                   "topic": str(t.get("topic", "") or "").strip()}
                  for i, t in enumerate(s.get("teams", [])) if isinstance(t, dict)]
    s["judges"] = [{"dept": str(j.get("dept", "")).strip() or f"심사위원 {i + 1}"}
                   for i, j in enumerate(s.get("judges", [])) if isinstance(j, dict)]
    s["criteria"] = [{"name": str(c.get("name", "")).strip() or f"항목 {i + 1}",
                      "max": max(1.0, _num(c.get("max"), 10.0))}
                     for i, c in enumerate(s.get("criteria", [])) if isinstance(c, dict)]
    return s


def get_settings() -> dict:
    with _lock:
        c = _connect()
        try:
            return _get(c, "settings")
        finally:
            c.close()


def _structure(s):
    return (len(s["teams"]), len(s["judges"]), tuple((c["max"]) for c in s["criteria"]))


def settings_impact(new: dict) -> dict:
    new = normalize_settings(new)
    old = get_settings()
    n_scores = _read_one("SELECT COUNT(*) AS n FROM scores")["n"]
    return {"structural": _structure(old) != _structure(new), "scores": n_scores,
            "judges_changed": len(old["judges"]) != len(new["judges"])}


def save_settings(new: dict, reset_if_needed: bool = False) -> dict:
    """설정 저장. 팀/심사위원 수나 배점 구조가 바뀌고 점수가 있으면 reset_if_needed=True가 필요합니다."""
    new = normalize_settings(new)
    if len(new["teams"]) < 2:
        raise FlowError("발표팀은 2팀 이상이어야 합니다.")
    if len(new["judges"]) < 1 or len(new["criteria"]) < 1:
        raise FlowError("심사위원과 평가 항목은 1개 이상이어야 합니다.")
    with _tx() as c:
        old = _get(c, "settings")
        structural = _structure(old) != _structure(new)
        has_scores = c.execute("SELECT COUNT(*) AS n FROM scores").fetchone()["n"] > 0
        if structural:
            if has_scores and not reset_if_needed:
                raise FlowError("이미 입력된 점수가 있어 구조를 바꾸려면 초기화가 필요합니다.")
            c.execute("DELETE FROM scores")
            _put(c, "state", copy.deepcopy(DEFAULT_STATE))
            if len(old["judges"]) != len(new["judges"]):
                _make_seats(c, len(new["judges"]))
        _put(c, "settings", new)
    return new


# ───────────────────────── 읽기 도우미
def _read(sql, args=()):
    with _lock:
        c = _connect()
        try:
            return c.execute(sql, args).fetchall()
        finally:
            c.close()


def _read_one(sql, args=()):
    rows = _read(sql, args)
    return rows[0] if rows else None


def get_state() -> dict:
    with _lock:
        c = _connect()
        try:
            st = _get(c, "state", copy.deepcopy(DEFAULT_STATE))
        finally:
            c.close()
    for k, v in DEFAULT_STATE.items():
        st.setdefault(k, copy.deepcopy(v))
    return st


def _norm(dept: str) -> str:
    return "".join(str(dept or "").split())


def eligible(S, judge_idx: int, team_idx: int) -> bool:
    """본인 관실 팀이면 평가 대상에서 제외합니다."""
    jd, td = _norm(S["judges"][judge_idx]["dept"]), _norm(S["teams"][team_idx]["dept"])
    return not (jd and td and jd == td)


def eligible_judges(S, team_idx: int) -> list[int]:
    return [j for j in range(len(S["judges"])) if eligible(S, j, team_idx)]


def current_team(st) -> int | None:
    if st["phase"] == "running" and st["order"] and 0 <= st["cur"] < len(st["order"]):
        return st["order"][st["cur"]]
    return None


# ───────────────────────── 심사위원 자리
def seats(now: float | None = None) -> list[dict]:
    now = now or time.time()
    S = get_settings()
    out = []
    for r in _read("SELECT idx, token, claimed_at, last_seen FROM seats ORDER BY idx"):
        if r["idx"] >= len(S["judges"]):
            continue
        out.append({"idx": r["idx"], "token": r["token"], "dept": S["judges"][r["idx"]]["dept"],
                    "label": _label(r["idx"]), "claimed": r["claimed_at"] is not None,
                    "online": bool(r["last_seen"] and now - r["last_seen"] < ONLINE_SEC)})
    return out


def _label(i: int) -> str:
    return chr(65 + i) if i < 26 else f"J{i + 1}"


def judge_by_token(token: str | None) -> dict | None:
    if not token:
        return None
    r = _read_one("SELECT idx, token FROM seats WHERE token=?", (str(token),))
    if not r:
        return None
    S = get_settings()
    if r["idx"] >= len(S["judges"]):
        return None
    return {"idx": r["idx"], "token": r["token"], "dept": S["judges"][r["idx"]]["dept"], "label": _label(r["idx"])}


def claim(idx: int) -> str:
    now = time.time()
    with _tx() as c:
        r = c.execute("SELECT token FROM seats WHERE idx=?", (int(idx),)).fetchone()
        if not r:
            raise FlowError("없는 심사위원 자리입니다.")
        c.execute("UPDATE seats SET claimed_at=?, last_seen=? WHERE idx=?", (now, now, int(idx)))
        return r["token"]


def touch(token: str):
    with _tx() as c:
        c.execute("UPDATE seats SET last_seen=?, claimed_at=COALESCE(claimed_at, ?) WHERE token=?",
                  (time.time(), time.time(), str(token)))


def release_seat(idx: int):
    with _tx() as c:
        c.execute("UPDATE seats SET claimed_at=NULL, last_seen=NULL WHERE idx=?", (int(idx),))


# ───────────────────────── 점수
def total_of(vals) -> float:
    return round(sum(float(v) for v in vals), 4)


def submit(token: str, team: int, vals: list, proxy: bool = False) -> float:
    """점수 제출(같은 팀에 다시 내면 덮어씀). proxy=True는 관리자 대리 입력으로 마감 후에도 허용합니다."""
    with _tx() as c:
        S, st = _get(c, "settings"), _get(c, "state")
        seat = c.execute("SELECT idx FROM seats WHERE token=?", (str(token),)).fetchone()
        if not seat or seat["idx"] >= len(S["judges"]):
            raise FlowError("심사위원 정보를 찾을 수 없습니다. QR로 다시 접속해 주세요.")
        if not (0 <= int(team) < len(S["teams"])):
            raise FlowError("없는 팀입니다.")
        if not eligible(S, seat["idx"], int(team)):
            raise FlowError("본인 관실 팀은 평가하지 않습니다.")
        if not proxy:
            if current_team(st) != int(team):
                raise FlowError("지금은 이 팀의 평가 시간이 아닙니다.")
            if int(team) in st["closed"]:
                raise FlowError("이미 마감된 평가입니다.")
        if len(vals) != len(S["criteria"]):
            raise FlowError("모든 항목을 입력해 주세요.")
        clean = []
        for v, crit in zip(vals, S["criteria"]):
            if v is None:
                raise FlowError(f"'{crit['name']}' 점수를 선택해 주세요.")
            v = float(v)
            if not (0 <= v <= crit["max"] + 1e-9):
                raise FlowError(f"'{crit['name']}' 점수는 0~{fmt(crit['max'])}점 사이여야 합니다.")
            clean.append(round(v, 4))
        tot = total_of(clean)
        c.execute("INSERT INTO scores(token, team, vals, total, proxy, updated_at) VALUES(?,?,?,?,?,?)"
                  " ON CONFLICT(token, team) DO UPDATE SET vals=excluded.vals, total=excluded.total,"
                  " proxy=excluded.proxy, updated_at=excluded.updated_at",
                  (str(token), int(team), json.dumps(clean), tot, 1 if proxy else 0, time.time()))
        return tot


def my_score(token: str, team: int) -> dict | None:
    r = _read_one("SELECT vals, total, proxy FROM scores WHERE token=? AND team=?", (str(token), int(team)))
    return {"vals": json.loads(r["vals"]), "total": r["total"], "proxy": bool(r["proxy"])} if r else None


def delete_score(token: str, team: int):
    with _tx() as c:
        c.execute("DELETE FROM scores WHERE token=? AND team=?", (str(token), int(team)))


def submission_status(team: int) -> dict:
    """팀별 제출 현황: 몇 명 중 몇 명, 미제출 관실 목록(관리자용)."""
    S = get_settings()
    tok = {r["idx"]: r["token"] for r in _read("SELECT idx, token FROM seats")}
    done = {r["token"] for r in _read("SELECT token FROM scores WHERE team=?", (int(team),))}
    el = eligible_judges(S, team)
    missing = [S["judges"][j]["dept"] for j in el if tok.get(j) not in done]
    return {"eligible": len(el), "submitted": len(el) - len(missing), "missing": missing}


def fmt(x) -> str:
    x = float(x)
    return str(int(round(x))) if abs(x - round(x)) < 1e-9 else f"{x:.2f}".rstrip("0").rstrip(".")


# ───────────────────────── 집계
def ranking() -> dict:
    """팀별 반영 평균(최고·최저 각 1개 제외)과 순위. 동점은 전체 평균으로 가르고, 그래도 같으면 관리자 결정."""
    S, st = get_settings(), get_state()
    tok2idx = {r["token"]: r["idx"] for r in _read("SELECT idx, token FROM seats")}
    by_team = {t: [] for t in range(len(S["teams"]))}
    for r in _read("SELECT token, team, vals, total FROM scores"):
        j = tok2idx.get(r["token"])
        t = r["team"]
        if j is None or t not in by_team or j >= len(S["judges"]) or not eligible(S, j, t):
            continue
        by_team[t].append((r["total"], json.loads(r["vals"])))
    rows = []
    for t, items in by_team.items():
        totals = sorted((x[0] for x in items), reverse=True)
        n = len(totals)
        trimmed = S["trim_high_low"] and n >= 3
        kept = totals[1:-1] if trimmed else totals
        crit_avg = [round(mean(v[1][k] for v in items), 2) if items else None for k in range(len(S["criteria"]))]
        rows.append({
            "team": t, "name": S["teams"][t]["name"], "dept": S["teams"][t]["dept"], "topic": S["teams"][t]["topic"],
            "order_no": (st["order"].index(t) + 1) if t in st["order"] else None,
            "n": n, "eligible": len(eligible_judges(S, t)), "totals": totals,
            "dropped_high": totals[0] if trimmed else None, "dropped_low": totals[-1] if trimmed else None,
            "kept": kept, "score": round(mean(kept), 4) if kept else None,
            "raw_mean": round(mean(totals), 4) if totals else None, "crit_avg": crit_avg,
        })

    def key(r):
        return (r["score"] is None, -(r["score"] or 0), -(r["raw_mean"] or 0))

    rows.sort(key=lambda r: (key(r), r["order_no"] or 99, r["team"]))
    rank, prev = 0, None
    for i, r in enumerate(rows):
        k = key(r)
        if k != prev:
            rank, prev = i + 1, k
        r["rank"] = rank if r["score"] is not None else None
    top = [r["team"] for r in rows if r["rank"] == 1]
    override = st.get("winner_override")
    if override is not None and override in top:
        winner = override
    elif len(top) == 1:
        winner = top[0]
    else:
        winner = None
    return {"rows": rows, "top_tied": top if len(top) > 1 else [], "winner": winner}


# ───────────────────────── 진행 단계
def _state_tx():
    return _tx()


def _set_state(c, st):
    _put(c, "state", st)


def draw_order() -> list[int]:
    with _tx() as c:
        S, st = _get(c, "settings"), _get(c, "state")
        if st["phase"] not in ("lobby", "order"):
            raise FlowError("발표가 시작된 뒤에는 순서를 다시 뽑을 수 없습니다.")
        order = list(range(len(S["teams"])))
        secrets.SystemRandom().shuffle(order)
        st.update(phase="order", order=order, drawn_at=time.time(), cur=0, step="presenting", closed=[])
        _set_state(c, st)
        return order


def start():
    with _tx() as c:
        st = _get(c, "state")
        if st["phase"] != "order":
            raise FlowError("먼저 발표 순서를 추첨해 주세요.")
        st.update(phase="running", cur=0, step="presenting", closed=[])
        _set_state(c, st)


def request_scoring():
    with _tx() as c:
        st = _get(c, "state")
        if st["phase"] != "running":
            raise FlowError("발표 진행 중이 아닙니다.")
        st["step"] = "scoring"
        _set_state(c, st)


def close_current():
    """현재 팀 평가를 마감하고 다음 발표로 넘어갑니다. 마지막 팀이면 집계 단계로."""
    with _tx() as c:
        st = _get(c, "state")
        if st["phase"] != "running":
            raise FlowError("발표 진행 중이 아닙니다.")
        team = st["order"][st["cur"]]
        if team not in st["closed"]:
            st["closed"].append(team)
        if st["cur"] < len(st["order"]) - 1:
            st["cur"] += 1
            st["step"] = "presenting"
        else:
            st["phase"] = "done"
        _set_state(c, st)


def back():
    """한 단계 되돌리기 (실수로 넘겼을 때)."""
    with _tx() as c:
        st = _get(c, "state")
        ph = st["phase"]
        if ph == "reveal":
            st["phase"] = "done"
        elif ph == "done":
            st["phase"] = "running"
            st["cur"] = len(st["order"]) - 1
            st["step"] = "scoring"
            last = st["order"][-1]
            st["closed"] = [t for t in st["closed"] if t != last]
        elif ph == "running":
            if st["step"] == "scoring":
                st["step"] = "presenting"
            elif st["cur"] > 0:
                st["cur"] -= 1
                prev = st["order"][st["cur"]]
                st["closed"] = [t for t in st["closed"] if t != prev]
                st["step"] = "scoring"
            else:
                st["phase"] = "order"
        elif ph == "order":
            if c.execute("SELECT COUNT(*) AS n FROM scores").fetchone()["n"]:
                raise FlowError("입력된 점수가 있어 추첨 전으로 돌아갈 수 없습니다.")
            st.update(phase="lobby", order=[], drawn_at=0.0)
        _set_state(c, st)


def set_winner_override(team: int | None):
    with _tx() as c:
        st = _get(c, "state")
        st["winner_override"] = None if team is None else int(team)
        _set_state(c, st)


def reveal():
    r = ranking()
    if r["winner"] is None:
        raise FlowError("최우수 팀이 정해지지 않았습니다. 동점이면 먼저 최우수 팀을 선택해 주세요."
                        if r["top_tied"] else "집계할 점수가 없습니다.")
    with _tx() as c:
        st = _get(c, "state")
        if st["phase"] not in ("done", "reveal"):
            raise FlowError("모든 발표의 평가가 마감된 뒤에 발표할 수 있습니다.")
        st.update(phase="reveal", revealed_at=time.time())
        _set_state(c, st)


def reset_progress():
    """점수와 진행 상황만 지웁니다(설정·심사위원 접속은 유지)."""
    with _tx() as c:
        c.execute("DELETE FROM scores")
        _put(c, "state", copy.deepcopy(DEFAULT_STATE))


def reset_seats():
    """모든 심사위원 접속 정보를 새로 발급합니다(전원 다시 QR 접속 필요)."""
    with _tx() as c:
        _make_seats(c, len(_get(c, "settings")["judges"]))


def factory_reset():
    """settings.json 기준으로 전부 처음 상태로."""
    s = load_file_settings()
    with _tx() as c:
        c.execute("DELETE FROM scores")
        _put(c, "settings", s)
        _put(c, "state", copy.deepcopy(DEFAULT_STATE))
        _make_seats(c, len(s["judges"]))


# ───────────────────────── 화면용 요약
def snapshot(now: float | None = None) -> dict:
    now = now or time.time()
    S, st = get_settings(), get_state()
    se = seats(now)
    cur = current_team(st)
    sub = submission_status(cur) if cur is not None else None
    data = {"S": S, "st": st, "now": now, "judges": len(S["judges"]),
            "online": sum(1 for s in se if s["online"]), "claimed": sum(1 for s in se if s["claimed"]),
            "seats": se, "cur": cur, "sub": sub, "winner": None, "top_tied": []}
    if st["phase"] in ("done", "reveal"):
        r = ranking()
        data["winner"], data["top_tied"] = r["winner"], r["top_tied"]
    return data


# ───────────────────────── 내보내기 / 백업
def export_excel() -> bytes:
    """결재·기록용 엑셀. 심사위원별 점수는 남기지 않고 팀별로 받은 점수만 익명으로 정리합니다."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill

    S, st = get_settings(), get_state()
    R = ranking()
    awards = {}
    for r in R["rows"]:
        if r["score"] is not None:
            awards[r["team"]] = "최우수" if r["team"] == R["winner"] else "우수"
    res = pd.DataFrame([{
        "순위": r["rank"], "팀": r["name"], "소속": r["dept"], "발표 순서": r["order_no"],
        "평가 인원": f'{r["n"]} / {r["eligible"]}',
        "반영 평균(최고·최저 제외)": r["score"], "전체 평균": r["raw_mean"],
        "제외된 최고점": r["dropped_high"], "제외된 최저점": r["dropped_low"],
        "수상": awards.get(r["team"], ""),
    } for r in R["rows"]])
    long = []
    for r in sorted(R["rows"], key=lambda x: (x["order_no"] or 99)):
        for i, tot in enumerate(r["totals"]):
            note = ""
            if r["dropped_high"] is not None and i == 0:
                note = "최고점 제외"
            elif r["dropped_low"] is not None and i == len(r["totals"]) - 1:
                note = "최저점 제외"
            long.append({"팀": r["name"], "점수(높은 순)": tot, "비고": note})
    anon = pd.DataFrame(long or [{"팀": "", "점수(높은 순)": None, "비고": ""}])
    crit = pd.DataFrame([dict({"팀": r["name"]}, **{
        f'{c["name"]} ({fmt(c["max"])}점)': r["crit_avg"][k] for k, c in enumerate(S["criteria"])})
        for r in sorted(R["rows"], key=lambda x: (x["order_no"] or 99))])
    info = pd.DataFrame([
        {"항목": "행사", "내용": f'{S["title_top"]} {S["title_main"]}'.strip()},
        {"항목": "일시·장소", "내용": S["date_line"]},
        {"항목": "심사위원", "내용": f'{len(S["judges"])}명 (본인 관실 팀은 평가 제외)'},
        {"항목": "평가 방식", "내용": (S["guide_note"] + " " if S["guide_note"] else "")
                                 + " · ".join(f'{c["name"]} {fmt(c["max"])}점' for c in S["criteria"])
                                 + (" — 평가기준: " + ", ".join(f'{g["name"]}({g["weight"]})' for g in S["guide"])
                                    if S["guide"] else "")},
        {"항목": "집계 기준", "내용": ("팀별로 받은 총점에서 최고점·최저점 각 1개를 제외한 평균"
                                  if S["trim_high_low"] else "팀별로 받은 총점의 평균")
                                 + ", 동점이면 전체 평균, 그래도 같으면 심사위원장 결정"},
        {"항목": "최우수", "내용": next((r["name"] for r in R["rows"] if r["team"] == R["winner"]), "미정")},
        {"항목": "익명 처리", "내용": "심사위원별 점수는 기록하지 않고 팀별로 받은 점수만 정리"},
        {"항목": "생성 시각", "내용": time.strftime("%Y-%m-%d %H:%M", time.localtime())},
    ])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in (("결과", res), ("팀별 받은 점수(익명)", anon), ("항목별 평균", crit), ("개요", info)):
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A8A")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width * 1.6 + 2))
            ws.freeze_panes = "A2"
    return buf.getvalue()


def backup() -> bytes:
    with _lock:
        c = _connect()
        try:
            data = {
                "settings": _get(c, "settings"), "state": _get(c, "state"),
                "seats": [dict(r) for r in c.execute("SELECT * FROM seats")],
                "scores": [dict(r) for r in c.execute("SELECT * FROM scores")],
                "saved_at": time.time(),
            }
        finally:
            c.close()
    return json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")


def restore(raw: bytes):
    data = json.loads(raw.decode("utf-8"))
    for k in ("settings", "state", "seats", "scores"):
        if k not in data:
            raise FlowError("백업 파일 형식이 아닙니다.")
    with _tx() as c:
        _put(c, "settings", normalize_settings(data["settings"]))
        _put(c, "state", data["state"])
        c.execute("DELETE FROM seats")
        for r in data["seats"]:
            c.execute("INSERT INTO seats(idx, token, claimed_at, last_seen) VALUES(?,?,?,?)",
                      (r["idx"], r["token"], r.get("claimed_at"), r.get("last_seen")))
        c.execute("DELETE FROM scores")
        for r in data["scores"]:
            c.execute("INSERT INTO scores(token, team, vals, total, proxy, updated_at) VALUES(?,?,?,?,?,?)",
                      (r["token"], r["team"], r["vals"], r["total"], r.get("proxy", 0), r["updated_at"]))
